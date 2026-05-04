import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

pptx_path = r'C:\Users\AkashK\Downloads\Earth-Watch\full.pptx'
prs = Presentation(pptx_path)

# Mining slides = 9, 10, 11 (0-indexed: 8, 9, 10)
for si in [8, 9, 10]:
    slide = prs.slides[si]
    print(f'\n{"="*70}')
    print(f'SLIDE {si+1}')
    print(f'{"="*70}')
    
    for j, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            if shape.shape_type == 13:
                print(f'  [{j}] IMAGE L={shape.left} T={shape.top} W={shape.width} H={shape.height}')
            continue
        
        full_text = []
        for p in shape.text_frame.paragraphs:
            t = p.text.strip()
            if t:
                full_text.append(t)
        
        if full_text:
            combined = ' | '.join(full_text)[:120]
            # font from first non-empty run
            finfo = ''
            for p in shape.text_frame.paragraphs:
                if p.runs:
                    r = p.runs[0]
                    sz = int(r.font.size / 12700) if r.font.size else 0
                    b = 'B' if r.font.bold else ''
                    try:
                        c = f'#{r.font.color.rgb}' if r.font.color.type else ''
                    except:
                        c = ''
                    finfo = f'{sz}pt {b} {c}'
                    break
            print(f'  [{j}] "{combined}" [{finfo}]')
